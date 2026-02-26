"""
Команда импорта данных из xlsx файлов в PostgreSQL.
Запуск: python manage.py import_data
Файлы xlsx и изображения должны лежать в папке import/ в корне проекта.
"""
import os
import shutil
from datetime import datetime
from django.core.management.base import BaseCommand
from django.conf import settings
import openpyxl

from store.models import (
    Role, User, Category, Manufacturer, Supplier,
    Unit, Product, PickupPoint, Order, OrderItem
)


class Command(BaseCommand):
    help = 'Импорт данных из xlsx файлов'

    def add_arguments(self, parser):
        parser.add_argument(
            '--path', type=str, default='import',
            help='Путь к папке с файлами импорта (по умолчанию: import/)'
        )

    def handle(self, *args, **options):
        base_path = options['path']

        # 1. Создаём роли
        self.stdout.write('Создание ролей...')
        role_admin, _ = Role.objects.get_or_create(role_name='Администратор')
        role_manager, _ = Role.objects.get_or_create(role_name='Менеджер')
        role_client, _ = Role.objects.get_or_create(role_name='Авторизированный клиент')

        # 2. Импорт пользователей
        self.stdout.write('Импорт пользователей...')
        user_file = os.path.join(base_path, 'user_import.xlsx')
        wb = openpyxl.load_workbook(user_file)
        ws = wb.active
        user_map = {}  # ФИО -> User

        for row in ws.iter_rows(min_row=2, values_only=True):
            role_name, fio, login, password = row[0], row[1], row[2], row[3]
            if not fio or not login:
                continue

            # Парсим ФИО
            parts = fio.strip().split()
            last_name = parts[0] if len(parts) > 0 else ''
            first_name = parts[1] if len(parts) > 1 else ''
            patronymic = parts[2] if len(parts) > 2 else ''

            # Определяем роль
            role = role_client
            if 'Администратор' in (role_name or ''):
                role = role_admin
            elif 'Менеджер' in (role_name or ''):
                role = role_manager

            if not User.objects.filter(username=login).exists():
                user = User.objects.create_user(
                    username=login,
                    password=password,
                    first_name=first_name,
                    last_name=last_name,
                    patronymic=patronymic,
                    role=role,
                )
                self.stdout.write(f'  + {fio} ({role_name})')
            else:
                user = User.objects.get(username=login)

            user_map[fio.strip()] = user

        # 3. Импорт товаров
        self.stdout.write('Импорт товаров...')
        tovar_file = os.path.join(base_path, 'Tovar.xlsx')
        wb = openpyxl.load_workbook(tovar_file)
        ws = wb.active
        product_map = {}  # артикул -> Product

        for row in ws.iter_rows(min_row=2, values_only=True):
            article = row[0]
            if not article:
                continue

            name = row[1] or ''
            unit_name = row[2] or 'шт.'
            price = row[3] or 0
            supplier_name = row[4] or ''
            manufacturer_name = row[5] or ''
            category_name = row[6] or ''
            discount = row[7] or 0
            qty = row[8] or 0
            description = (row[9] or '').replace('\xa0', ' ').strip()
            photo_filename = row[10]

            # get_or_create для справочников
            category, _ = Category.objects.get_or_create(category_name=category_name)
            manufacturer, _ = Manufacturer.objects.get_or_create(manufacturer_name=manufacturer_name)
            supplier, _ = Supplier.objects.get_or_create(supplier_name=supplier_name)
            unit, _ = Unit.objects.get_or_create(unit_name=unit_name)

            product, created = Product.objects.get_or_create(
                article=str(article).strip(),
                defaults={
                    'product_name': name,
                    'unit': unit,
                    'price': price,
                    'supplier': supplier,
                    'manufacturer': manufacturer,
                    'category': category,
                    'discount': discount,
                    'quantity_in_stock': qty,
                    'description': description,
                }
            )

            # Копируем фото в media/products/
            if photo_filename:
                src = os.path.join(base_path, photo_filename)
                if os.path.exists(src):
                    dst_dir = os.path.join(settings.MEDIA_ROOT, 'products')
                    os.makedirs(dst_dir, exist_ok=True)
                    dst = os.path.join(dst_dir, photo_filename)
                    shutil.copy2(src, dst)
                    product.photo = f'products/{photo_filename}'
                    product.save()

            product_map[str(article).strip()] = product
            if created:
                self.stdout.write(f'  + {article} — {name}')

        # 4. Импорт пунктов выдачи
        self.stdout.write('Импорт пунктов выдачи...')
        pp_file = os.path.join(base_path, 'Пункты выдачи_import.xlsx')
        wb = openpyxl.load_workbook(pp_file)
        ws = wb.active
        pp_map = {}  # порядковый номер -> PickupPoint

        for i, row in enumerate(ws.iter_rows(min_row=1, values_only=True), start=1):
            address = (row[0] or '').replace('\xa0', ' ').strip()
            if not address:
                continue
            pp, _ = PickupPoint.objects.get_or_create(address=address)
            pp_map[i] = pp

        # 5. Импорт заказов
        self.stdout.write('Импорт заказов...')
        order_file = os.path.join(base_path, 'Заказ_import.xlsx')
        wb = openpyxl.load_workbook(order_file)
        ws = wb.active

        for row in ws.iter_rows(min_row=2, values_only=True):
            order_num = row[0]
            articles_str = row[1]
            order_date_raw = row[2]
            delivery_date_raw = row[3]
            pp_index = row[4]
            client_fio = (row[5] or '').strip()
            pickup_code = row[6]
            status = (row[7] or 'Новый').strip()

            if not order_num or not articles_str:
                continue

            # Парсим дату заказа
            order_date = self._parse_date(order_date_raw)
            delivery_date = self._parse_date(delivery_date_raw)
            if not order_date:
                self.stdout.write(self.style.WARNING(
                    f'  ! Заказ {order_num}: некорректная дата "{order_date_raw}", пропуск'
                ))
                continue

            # Находим пункт выдачи по номеру
            pp = pp_map.get(int(pp_index)) if pp_index else None
            if not pp:
                self.stdout.write(self.style.WARNING(
                    f'  ! Заказ {order_num}: пункт выдачи {pp_index} не найден, пропуск'
                ))
                continue

            # Находим клиента
            client = user_map.get(client_fio)
            if not client:
                self.stdout.write(self.style.WARNING(
                    f'  ! Заказ {order_num}: клиент "{client_fio}" не найден, пропуск'
                ))
                continue

            order = Order.objects.create(
                order_date=order_date,
                delivery_date=delivery_date,
                pickup_point=pp,
                client=client,
                pickup_code=pickup_code or 0,
                status=status,
            )

            # Парсим состав: "А112Т4, 2, F635R4, 2" -> [(А112Т4, 2), (F635R4, 2)]
            parts = [p.strip() for p in str(articles_str).split(',')]
            for j in range(0, len(parts) - 1, 2):
                art = parts[j].strip()
                try:
                    qty = int(parts[j + 1].strip())
                except (ValueError, IndexError):
                    qty = 1

                product = product_map.get(art)
                if product:
                    OrderItem.objects.create(order=order, product=product, quantity=qty)
                else:
                    self.stdout.write(self.style.WARNING(
                        f'    ! Товар {art} не найден для заказа {order_num}'
                    ))

            self.stdout.write(f'  + Заказ №{order_num}')

        self.stdout.write(self.style.SUCCESS('Импорт завершён!'))

    def _parse_date(self, raw):
        """Парсим дату из xlsx (может быть datetime или строка)"""
        if isinstance(raw, datetime):
            return raw.date()
        if isinstance(raw, str):
            for fmt in ('%Y-%m-%d', '%d.%m.%Y'):
                try:
                    return datetime.strptime(raw.strip(), fmt).date()
                except ValueError:
                    continue
        return None
